# Builds proposal-slides.pptx — a 10-slide, fully editable FYP proposal-defense deck for Metamorph.
# Mirrors proposal-slides.html. Run:  python data_work/build_pptx.py
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "proposal-slides.pptx")
SHOT_TESTING = os.path.join(ROOT, "scripts", "d-testing.png")
SHOT_REPORT = os.path.join(ROOT, "scripts", "d-done.png")

# ---- palette ----
INK   = RGBColor(0x1A, 0x1F, 0x26)
MUTE  = RGBColor(0x5B, 0x66, 0x75)
FAINT = RGBColor(0x8B, 0x94, 0xA1)
LINE  = RGBColor(0xDD, 0xE1, 0xE6)
PANEL = RGBColor(0xF5, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACC   = RGBColor(0x2F, 0x6B, 0xFF)   # testing / execution
ACC_BG= RGBColor(0xEA, 0xF0, 0xFF)
AGENT = RGBColor(0x6A, 0x5A, 0xCD)   # llm / agent
AGENT_BG = RGBColor(0xEE, 0xEC, 0xFB)
PASS_ = RGBColor(0x1F, 0x9D, 0x63)
FAIL_ = RGBColor(0xD6, 0x32, 0x3A)
WARN_ = RGBColor(0xB9, 0x79, 0x1A)
PLACEHOLDER = RGBColor(0x2F, 0x6B, 0xFF)

F_TITLE = "Segoe UI Semibold"
F_BODY  = "Segoe UI"
F_MONO  = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

MX = Inches(0.62)          # outer margin
CW = Inches(13.333 - 1.24) # content width


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    return s


def _noshadow(sh):
    try:
        sh.shadow.inherit = False
    except Exception:
        pass


def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    _noshadow(sh)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, runs, size=13, color=INK, font=F_BODY, bold=False, space_after=6,
         align=PP_ALIGN.LEFT, bullet=False, line=1.15, first=False):
    """runs: str, or list of (text, {opts}) tuples for inline styling."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, o in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(o.get("size", size))
        r.font.name = o.get("font", font)
        r.font.bold = o.get("bold", bold)
        r.font.italic = o.get("italic", False)
        r.font.color.rgb = o.get("color", color)
    _set_bullet(p, bullet)
    return p


def _set_bullet(p, on):
    pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    if on:
        pPr.set("indent", str(-Inches(0.18)))
        pPr.set("marL", str(Inches(0.18)))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
    else:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def header(s, n, title):
    tf = textbox(s, MX, Inches(0.42), CW - Inches(1.2), Inches(0.7))
    para(tf, title, size=25, color=INK, font=F_TITLE, bold=True, space_after=0, first=True)
    tf2 = textbox(s, prs.slide_width - Inches(1.7), Inches(0.5), Inches(1.1), Inches(0.4))
    para(tf2, f"{n:02d} / 10", size=11, color=FAINT, font=F_MONO, align=PP_ALIGN.RIGHT,
         space_after=0, first=True)
    rect(s, MX, Inches(1.22), CW, Pt(1.4), fill=LINE)


def panel(s, x, y, w, h, title=None, fill=PANEL, line=LINE, title_color=FAINT):
    rect(s, x, y, w, h, fill=fill, line=line, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.045)
    pad = Inches(0.22)
    tf = textbox(s, x + pad, y + pad, w - 2 * pad, h - 2 * pad)
    if title:
        para(tf, title.upper(), size=10.5, color=title_color, font=F_MONO, bold=True,
             space_after=9, first=True)
    return tf


def kicker(s, x, y, w, text):
    tf = textbox(s, x, y, w, Inches(0.3))
    para(tf, text.upper(), size=10.5, color=FAINT, font=F_MONO, bold=True, space_after=0, first=True)


def bullets(tf, items, size=12.5, gap=7):
    """items: (text | list-of-runs, {color, agent})"""
    first_done = len(tf.paragraphs[0].runs) > 0 or (tf.paragraphs[0].text != "")
    for it in items:
        content, o = it if isinstance(it, tuple) else (it, {})
        col = AGENT if o.get("agent") else o.get("color", INK)
        if isinstance(content, str):
            content = [(content, {})]
        content = [(t, {**op, "color": op.get("color", col)}) for t, op in content]
        para(tf, content, size=size, space_after=gap, bullet=True, line=1.22,
             first=not first_done)
        first_done = True


def metric_tile(s, x, y, w, h, value, label, vcolor=INK):
    rect(s, x, y, w, h, fill=PANEL, line=LINE, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    tf = textbox(s, x + Inches(0.16), y + Inches(0.13), w - Inches(0.32), h - Inches(0.26))
    para(tf, value, size=22, color=vcolor, font=F_TITLE, bold=True, space_after=4, first=True)
    para(tf, label.upper(), size=9.5, color=FAINT, font=F_MONO, space_after=0)


PH = {"color": PLACEHOLDER, "italic": True}

# ================================================================= 1
s = slide()
kicker(s, MX, Inches(0.5), CW, "FYP Proposal Defense · FAST School of Computing")
tf2 = textbox(s, prs.slide_width - Inches(1.7), Inches(0.5), Inches(1.1), Inches(0.4))
para(tf2, "01 / 10", size=11, color=FAINT, font=F_MONO, align=PP_ALIGN.RIGHT, space_after=0, first=True)

tf = textbox(s, MX, Inches(1.9), CW, Inches(2.2))
para(tf, "Metamorph", size=64, color=INK, font=F_TITLE, bold=True, space_after=6, first=True)
para(tf, "An agentic metamorphic testing framework for computer vision models.",
     size=19, color=MUTE, space_after=0, line=1.25)

rect(s, MX, Inches(4.35), Pt(3), Inches(1.15), fill=ACC)
tf = textbox(s, MX + Inches(0.22), Inches(4.35), Inches(9.6), Inches(1.2))
para(tf, [
    ("We can measure a vision model's accuracy on its test set — but we have ", {}),
    ("no way to check it", {"bold": True}),
    (" against fog, night, motion blur, glare or occlusion, because those inputs have no labels and no test oracle.", {}),
], size=15.5, color=INK, line=1.4, space_after=0, first=True)

rect(s, MX, Inches(6.15), CW, Pt(1.2), fill=LINE)
tf = textbox(s, MX, Inches(6.35), CW, Inches(0.8))
para(tf, [
    ("Group   ", {"color": MUTE, "size": 12.5}),
    ("〈Member A〉 · 〈Member B〉 · 〈Member C〉", {**PH, "size": 12.5, "font": F_MONO}),
    ("        Supervisor   ", {"color": MUTE, "size": 12.5}),
    ("〈Supervisor〉", {**PH, "size": 12.5, "font": F_MONO}),
    ("        Category   ", {"color": MUTE, "size": 12.5}),
    ("R&D · ML testing", {"color": INK, "size": 12.5, "font": F_MONO}),
], space_after=0, first=True)

# ================================================================= 2
s = slide(); header(s, 2, "Problem and users")
kicker(s, MX, Inches(1.5), Inches(6.6), "Who has the problem")
tf = textbox(s, MX, Inches(1.85), Inches(6.6), Inches(1.6))
para(tf, "Any team shipping a trained detector, segmenter or tracker — autonomous driving, "
         "ADAS, agriculture, medical imaging, security — plus the ML researchers and QA "
         "engineers who must sign off on it.", size=15, color=INK, line=1.45, space_after=0, first=True)
kicker(s, MX, Inches(3.7), Inches(6.6), "Why it matters")
tf = textbox(s, MX, Inches(4.05), Inches(6.6), Inches(2.6))
para(tf, [
    ("The held-out test set never contains the weather, lighting and camera conditions the "
     "model meets in deployment. Models fail ", {}),
    ("silently", {"bold": True}),
    (" and the failures concentrate on vulnerable classes — pedestrians, cyclists — where an "
     "error is most costly. There is no standard method, no oracle, and no coverage measure "
     "for this today.", {}),
], size=13, color=MUTE, line=1.5, space_after=0, first=True)

tf = panel(s, Inches(7.7), Inches(1.5), Inches(5.0), Inches(4.4), "Current workflow")
bullets(tf, [
    "Collect & label a dataset",
    "Train the model",
    "Report accuracy / mAP on a held-out split",
    ([("Robustness?", {"bold": True}), (" checked ad-hoc — a few hand-picked images, "
      "hand-written augmentations, per person, per project", {})], {}),
    "Ship",
], size=12.5, gap=9)
para(tf, "The gap sits between step 3 and step 5.", size=11, color=FAINT, space_after=0)

# ================================================================= 3
s = slide(); header(s, 3, "Existing solutions and prior-FYP comparison")
rows = [
    ["Capability", "Adversarial testing\nFGSM / PGD", "Augmentation libs\nimgaug · Albumentations",
     "Existing MT for CV\nfixed 21 MRs", "Metamorph"],
    ["Oracle without new labels", "partial", "no", "yes — fixed", "yes — task-aware"],
    ["Realistic weather / lighting", "no", "pixel filters", "1 (sunny→snowy)", "generative edits"],
    ["Adapts to the model's domain", "no", "no", "no", "LLM-discovered MRs"],
    ["Checks meaning is preserved", "n/a", "no", "no", "CLIP gate"],
    ["Evidence-guided test selection", "no", "no", "no", "closed loop + memory"],
    ["Coverage & robustness report", "no", "no", "partial", "per condition & class"],
]
gt = s.shapes.add_table(len(rows), 5, MX, Inches(1.5), CW, Inches(3.9)).table
gt.columns[0].width = Inches(2.9)
for i in range(1, 5):
    gt.columns[i].width = Inches((12.09 - 2.9) / 4)
for r, row in enumerate(rows):
    gt.rows[r].height = Inches(0.52 if r else 0.62)
    for c, val in enumerate(row):
        cell = gt.cell(r, c)
        cell.margin_left = cell.margin_right = Inches(0.09)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = PANEL if c < 4 else ACC_BG
        else:
            cell.fill.fore_color.rgb = WHITE if c < 4 else ACC_BG
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run(); run.text = val
        run.font.name = F_MONO if r == 0 else F_BODY
        run.font.size = Pt(10 if r == 0 else 11)
        run.font.bold = (r == 0) or (c == 0) or (c == 4)
        if r == 0:
            run.font.color.rgb = FAINT
        elif c == 0:
            run.font.color.rgb = INK
        elif c == 4:
            run.font.color.rgb = INK
        else:
            run.font.color.rgb = {"no": FAIL_, "n/a": FAINT}.get(val, PASS_ if val.startswith("yes")
                                 else WARN_)
# strip default table style banding
tblPr = gt._tbl.find(qn("a:tblPr"))
if tblPr is not None:
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")

tf = textbox(s, MX, Inches(5.6), CW, Inches(1.4))
para(tf, [
    ("Prior FYP work   ", {"bold": True, "size": 12.5}),
    ("〈if any — e.g. \"Prior FYP applied the fixed 21 MRs to one classifier; no domain "
     "adaptation, no generative transforms, no memory.\"〉", {**PH, "size": 12}),
], line=1.45, space_after=6, first=True)
para(tf, [
    ("The gap:  ", {"bold": True, "color": ACC, "font": F_MONO, "size": 12.5}),
    ("no framework combines a label-free task-aware oracle with domain-adaptive MR discovery, "
     "generative transforms behind a semantic gate, and adaptive selection.", {"size": 12.5, "color": MUTE}),
], line=1.45, space_after=0)

# ================================================================= 4
s = slide(); header(s, 4, "R&D / recent technical work")
tf = panel(s, MX, Inches(1.5), Inches(6.05), Inches(5.5), "Sources we build on")
bullets(tf, [
    [("Metamorphic testing", {"bold": True}), (" — Chen et al., ACM CSUR 2018: test without "
      "an oracle via relations between inputs and outputs.", {})],
    [("MT for deep models", {"bold": True}), (" — DeepTest (ICSE 2018), DeepRoad (ASE 2018): "
      "GAN weather for driving models — but hand-fixed relations.", {})],
    [("The 21-MR library", {"bold": True}), (" for image classification — our general-purpose "
      "starting set.", {})],
    ([("InstructPix2Pix", {"bold": True}), (" — Brooks et al., CVPR 2023: edit an image from a "
      "text instruction.", {})], {"agent": True}),
    ([("CLIP", {"bold": True}), (" — Radford et al., ICML 2021: image↔text similarity, used "
      "here as the semantic gate.", {})], {"agent": True}),
    ([("Tool-using LLM agents", {"bold": True}), (" — plan-act-observe loops for MR discovery.", {})],
     {"agent": True}),
], size=12, gap=9)

tf = panel(s, Inches(7.0), Inches(1.5), Inches(5.7), Inches(5.5), "What we took from it",
           fill=ACC_BG, line=ACC, title_color=ACC)
bullets(tf, [
    "MT turns \"no labels\" into \"a checkable relation\" — the right primitive for deployment conditions.",
    [("Generative edits look real, ", {}), ("but can silently change the scene", {"bold": True}),
     (" → a semantic-preservation check is non-negotiable.", {})],
    [("A fixed MR set does not transfer across domains → discovery must be driven by a ", {}),
     ("domain description", {"bold": True}), (" and by ", {}), ("evidence collected so far", {"bold": True}),
     (".", {})],
    "The LLM must never decide pass/fail — deterministic transforms and objective oracles do.",
], size=12.5, gap=10)

# ================================================================= 5
s = slide(); header(s, 5, "Complex computing challenge")
tf = textbox(s, MX, Inches(1.5), Inches(7.4), Inches(5.5))
bullets(tf, [
    [("No ground truth for the oracle. ", {"bold": True}), ("\"Should this transform change the "
      "output, and how?\" must be synthesised per relation type — invariant, predictable, "
      "additive, relational — across classification, detection, segmentation and tracking.", {})],
    [("Open-ended relation space. ", {"bold": True}), ("The agent proposes new relations from a "
      "domain description and accumulated failures; each must pass schema, task-compatibility, "
      "domain, feasibility and semantic checks before it can run.", {})],
    [("Generative transforms break silently. ", {"bold": True}), ("An edit can add or remove "
      "objects — needs a CLIP gate with a regenerate-or-fallback loop.", {})],
    [("Search under budget. ", {"bold": True}), ("transform × image × parameter is a huge space; "
      "exhaustive testing is infeasible → adaptive, evidence-guided selection.", {})],
    [("Reproducible verdicts", {"bold": True}), (" from a pipeline that contains a stochastic LLM "
      "and a stochastic image editor.", {})],
], size=12.5, gap=11)

tf = panel(s, Inches(8.35), Inches(1.5), Inches(4.35), Inches(4.6), "CCP mapping")
bullets(tf, [
    "Depth of analysis & abstract reasoning",
    "No standard, off-the-shelf solution",
    "Novel integration: VLM + LLM + editor + CLIP + model + oracle + memory",
    [("Conflicting requirements: ", {}), ("realism ⟷ semantic fidelity ⟷ coverage ⟷ budget",
      {"font": F_MONO, "size": 11})],
], size=12, gap=9)
para(tf, "Systems, ML, program-analysis and HCI concerns in one loop.", size=11, color=FAINT, space_after=0)

# ================================================================= 6
s = slide(); header(s, 6, "Proposed solution and contribution")
# left: flow
tf = textbox(s, MX, Inches(1.45), Inches(6.3), Inches(0.55))
para(tf, [("Bootstrap sweep   ", {"bold": True, "font": F_MONO, "size": 12, "color": INK}),
          ("~10 images × transform library → testing memory", {"color": FAINT, "size": 10.5})],
     space_after=0, first=True)
steps = [
    ("1", "Pick image", "from the dataset", False),
    ("2", "Describe", "vision model → text caption", True),
    ("3", "Plan relation", "LLM + memory — may invent a new MR", True),
    ("4", "Transform", "InstructPix2Pix if generative · code otherwise", False),
    ("5", "Semantic check", "CLIP: is the meaning preserved?", False),
    ("6", "Run model", "only if the check passed", False),
    ("7", "Verify", "task-aware oracle — relation held?", False),
    ("8", "Update memory", "store image · transform · verdict", False),
]
rect(s, MX + Inches(0.12), Inches(2.05), Pt(2), Inches(3.55), fill=LINE)
y = Inches(2.05)
for num, name, sub, ag in steps:
    dot = rect(s, MX + Inches(0.02), y + Inches(0.06), Inches(0.18), Inches(0.18),
               fill=WHITE, line=(AGENT if ag else RGBColor(0xC7, 0xCD, 0xD6)), line_w=1.6,
               shape=MSO_SHAPE.OVAL)
    tf = textbox(s, MX + Inches(0.42), y, Inches(6.0), Inches(0.44))
    para(tf, [
        (f"{num}   ", {"font": F_MONO, "size": 10.5, "color": FAINT}),
        (name, {"bold": True, "size": 12.5, "color": (AGENT if ag else INK)}),
        (f"    {sub}", {"size": 11, "color": MUTE}),
    ], space_after=0, first=True)
    y += Inches(0.445)
tf = textbox(s, MX + Inches(0.42), y + Inches(0.05), Inches(6.0), Inches(0.7))
para(tf, [
    ("↻  ", {"bold": True, "color": AGENT}),
    ("memory feeds step 3 — the loop repeats under a fixed test budget, then compiles the ", {"size": 11, "color": FAINT}),
    ("robustness & coverage report", {"size": 11, "bold": True, "color": INK}),
    (".", {"size": 11, "color": FAINT}),
], line=1.4, space_after=0, first=True)

# right: modules + contribution
tf = panel(s, Inches(7.0), Inches(1.45), Inches(5.7), Inches(2.35), "Main modules")
para(tf, "Model adapter · Bootstrap sweep · MR Discovery Agent · MR Validator · Adaptive "
         "Selection Agent · Transformation Engine (InstructPix2Pix + deterministic) · Semantic "
         "Validator (CLIP) · Task-aware Oracle · Failure Analyzer · Orchestrator",
     size=11.5, color=MUTE, line=1.6, space_after=0)
tf = panel(s, Inches(7.0), Inches(4.0), Inches(5.7), Inches(3.0), "Contribution",
           fill=WHITE, line=ACC, title_color=ACC)
para(tf, [
    ("A domain-adaptive agentic MT framework that (1) ", {}),
    ("discovers and validates new metamorphic relations", {"bold": True}),
    (" from a plain-language domain description and accumulated failure evidence, (2) executes "
     "them with ", {}),
    ("realistic generative or deterministic transforms behind a CLIP semantic gate", {"bold": True}),
    (", (3) judges them with ", {}),
    ("task-aware oracles", {"bold": True}),
    (" across four CV tasks, and (4) ", {}),
    ("selects tests adaptively in a closed evidence-driven loop", {"bold": True}),
    (" — producing an interpretable robustness and coverage report.", {}),
], size=11.5, color=MUTE, line=1.5, space_after=0)

# ================================================================= 7
s = slide(); header(s, 7, "POC-lite / feasibility evidence")
s.shapes.add_picture(SHOT_TESTING, MX, Inches(1.5), width=Inches(6.15))
tf = textbox(s, MX, Inches(5.55), Inches(6.15), Inches(0.32))
para(tf, "PROTOTYPE — ONE ITERATION OF THE LOOP, LIVE", size=9.5, color=FAINT, font=F_MONO,
     space_after=0, first=True)
tf = textbox(s, MX, Inches(5.95), Inches(6.15), Inches(1.5))
para(tf, [
    ("10 real ", {}), ("KITTI", {"bold": True}),
    (" street scenes + ground-truth boxes from Hugging Face; deterministic + generative-style "
     "transforms (fog · rain · night · dusk · haze · glare · motion blur · occlusion · rotation "
     "· insertion) generated offline with detection overlays. The full pipeline runs end-to-end, "
     "including a CLIP regenerate case.", {}),
], size=11, color=MUTE, line=1.4, space_after=0, first=True)

gx = Inches(7.05)
panel(s, gx, Inches(1.5), Inches(5.65), Inches(3.0), "Risk experiment — what it found")
tw, th, gap = Inches(2.68), Inches(0.98), Inches(0.16)
ty = Inches(2.02)
metric_tile(s, gx + Inches(0.22), ty, tw, th, "0.91", "baseline mAP@0.5")
metric_tile(s, gx + Inches(0.22) + tw + gap, ty, tw, th, "5 / 10", "relations violated", FAIL_)
metric_tile(s, gx + Inches(0.22), ty + th + gap, tw, th, "pedestrian", "most affected class", FAIL_)
metric_tile(s, gx + Inches(0.22) + tw + gap, ty + th + gap, tw, th, "28%", "robustness · night & occl.", ACC)

tf = textbox(s, gx, Inches(4.9), Inches(5.65), Inches(1.3))
para(tf, [
    ("Violations concentrated under ", {}),
    ("fog, rain, night, occlusion and glare", {"bold": True}),
    (" — exactly the untested conditions — and produced a per-condition robustness profile.", {}),
], size=12, color=MUTE, line=1.5, space_after=10, first=True)
para(tf, [
    ("Full MVP not required", {"bold": True, "color": INK}),
    (" — this proves the workflow, the data path and that the hardest risks are addressable.",
     {"color": FAINT}),
], size=12, line=1.5, space_after=0)

# ================================================================= 8
s = slide(); header(s, 8, "Evaluation plan and risks")
tf = panel(s, MX, Inches(1.5), Inches(6.0), Inches(5.5), "How success is measured")
bullets(tf, [
    [("Oracle validity", {"bold": True}), (" — verdicts vs human judgement on a labelled sample   ", {}),
     ("≥ 90% agreement", {"font": F_MONO, "color": ACC, "size": 11})],
    [("MR discovery quality", {"bold": True}), (" — % of proposed relations that pass the validator and expert review", {})],
    [("Semantic gate", {"bold": True}), (" — CLIP precision/recall vs human \"did the meaning change?\" labels", {})],
    [("Fault-finding power", {"bold": True}), (" — real violations surfaced vs a manual audit and vs the fixed-21-MR baseline", {})],
    [("Coverage", {"bold": True}), (" — relation-type and condition coverage reached within budget", {})],
    [("Reproducibility", {"bold": True}), (" — identical inputs → identical verdicts", {})],
], size=12, gap=10)

rx = Inches(7.0)
risks = [
    ("Risk", "Mitigation"),
    ("LLM proposes invalid relations", "validator + human-in-the-loop gate"),
    ("Generative edit changes content", "CLIP gate → regenerate → deterministic fallback"),
    ("Compute cost of image editing", "cache · batch · budget · smaller editor"),
    ("Dataset licensing (KITTI = non-commercial)", "research use only; swap-in dataset interface"),
    ("Scope across 4 CV tasks", "detection first; others via the adapter"),
]
rt = s.shapes.add_table(len(risks), 2, rx, Inches(1.5), Inches(5.7), Inches(2.9)).table
rt.columns[0].width = Inches(2.75); rt.columns[1].width = Inches(2.95)
for r, (a, b) in enumerate(risks):
    rt.rows[r].height = Inches(0.42)
    for c, val in enumerate((a, b)):
        cell = rt.cell(r, c); cell.margin_left = cell.margin_right = Inches(0.09)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid(); cell.fill.fore_color.rgb = PANEL if r == 0 else WHITE
        p = cell.text_frame.paragraphs[0]; cell.text_frame.word_wrap = True
        run = p.add_run(); run.text = val
        run.font.name = F_MONO if r == 0 else F_BODY
        run.font.size = Pt(10 if r == 0 else 10.5)
        run.font.bold = (r == 0)
        run.font.color.rgb = FAINT if r == 0 else INK
tblPr = rt._tbl.find(qn("a:tblPr"))
if tblPr is not None:
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")

tf = panel(s, rx, Inches(4.65), Inches(5.7), Inches(2.35), "Tool / API / GenAI boundary",
           fill=AGENT_BG, line=AGENT, title_color=AGENT)
para(tf, [
    ("GenAI", {"bold": True}),
    (" is used to propose relations and edit images — every output is validated and can be "
     "rejected. ", {}),
    ("The oracle, the metrics and every pass/fail decision are deterministic code.", {"bold": True}),
    (" Third-party, all open / research-licensed: HF datasets, an LLM, InstructPix2Pix, CLIP. "
     "Our work: the agents, validators, oracles, orchestration and UI.", {}),
], size=11.5, color=MUTE, line=1.5, space_after=0)

# ================================================================= 9
s = slide(); header(s, 9, "Work division and iteration plan")
tf = panel(s, MX, Inches(1.5), Inches(6.4), Inches(5.5), "Ownership")
bullets(tf, [
    [("〈Member A〉", {**PH, "bold": True}), (" — model adapters · bootstrap sweep · deterministic "
      "transformation engine · detection oracle", {})],
    ([("〈Member B〉", {**PH, "bold": True}), (" — MR Discovery Agent · MR Validator · Adaptive "
      "Selection Agent · testing memory · LLM integration", {})], {"agent": True}),
    ([("〈Member C〉", {**PH, "bold": True}), (" — generative transform service (InstructPix2Pix) · "
      "CLIP semantic validator · failure analyzer · reporting UI", {})], {"agent": True}),
    [("Shared", {"bold": True}), (" — orchestrator · evaluation harness · report write-up", {})],
], size=12, gap=12)

tf = panel(s, Inches(7.35), Inches(1.5), Inches(5.35), Inches(4.2), "Iterations")
bullets(tf, [
    [("It 1 · FYP-1", {"bold": True}), (" — adapter + bootstrap sweep + deterministic transforms "
      "+ detection oracle + memory + demo on KITTI", {})],
    [("It 2 · FYP-1", {"bold": True}), (" — MR Discovery Agent + Validator + CLIP gate + "
      "rule-based adaptive selection + robustness report", {})],
    ([("It 3 · FYP-2", {"bold": True}), (" — generative transform service + regenerate loop + "
      "LLM planning from memory", {})], {"agent": True}),
    ([("It 4 · FYP-2", {"bold": True}), (" — segmentation & tracking oracles via the adapter + "
      "evaluation study + hardening", {})], {"agent": True}),
], size=11.5, gap=10)
tf = textbox(s, Inches(7.35), Inches(5.9), Inches(5.35), Inches(1.1))
para(tf, [
    ("Milestones", {"bold": True, "size": 11.5}),
    (" aligned to the handbook — Proposal Defense wk 4 · Mid wk 8 (It 1) · Poster wk 13 · Final "
     "FYP-1 wk 15 (It 2, ≥ 45% implementation).", {"size": 11.5, "color": MUTE}),
], line=1.45, space_after=0, first=True)

# ================================================================= 10
s = slide(); header(s, 10, "DEI, ethics and legal")
pw = Inches((12.09 - 0.8) / 3); px = MX
cols = [
    ("DEI", "The framework's purpose is equity in ML reliability. It deliberately surfaces "
     "failures on under-represented classes and conditions — pedestrians at night, darker skin "
     "tones under low light, regional weather — and reports per class and per condition, so "
     "disparities are visible instead of hidden in an average. Test scenes and domain "
     "descriptions span geographies, lighting and appearance."),
    ("Ethics", "No human-subjects data collection. Public research images with visible faces "
     "are handled per dataset terms (blurred on export where required). We never claim "
     "\"production ready\" — the report always states coverage and remaining budget, so results "
     "cannot be cherry-picked into a safety claim."),
    ("Legal", "Datasets under their licences — KITTI: CC BY-NC-SA 3.0, non-commercial research "
     "only. Model weights (CLIP, InstructPix2Pix) and the LLM under their own terms. All "
     "third-party components isolated behind interfaces; the framework, agents, oracles and UI "
     "are the group's own work. Testing memory stores only image ids, transform parameters and "
     "verdicts — no PII."),
]
for title, body in cols:
    tf = panel(s, px, Inches(1.55), pw, Inches(4.7), title)
    para(tf, body, size=12, color=MUTE, line=1.6, space_after=0)
    px += pw + Inches(0.4)

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
